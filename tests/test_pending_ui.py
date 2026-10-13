from models.schemas import StudyPackage
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pymupdf
from PIL import Image
import streamlit as st
from streamlit.testing.v1 import AppTest


ROOT = Path(__file__).resolve().parents[1]
DASHBOARD_PATH = ROOT / "pages" / "0_Dashboard.py"


def make_pdf(text: str) -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    file_bytes = document.tobytes()
    document.close()
    return file_bytes


def install_upload_mock(monkeypatch):
    upload = {"file": None}

    def fake_file_uploader(*_args, **_kwargs):
        return upload["file"]

    monkeypatch.setattr(st, "file_uploader", fake_file_uploader)
    return upload


def make_upload(filename: str, content: bytes):
    return SimpleNamespace(name=filename, getvalue=lambda: content)


def click_extract(app: AppTest) -> None:
    next(
        button for button in app.button if button.label == "Extract text"
    ).click().run()


def test_dashboard_create_course_button_toggles_setup_panel():
    app = AppTest.from_file(str(DASHBOARD_PATH)).run()

    assert not app.exception
    assert app.button(key="create_course_button").label == "Create course"
    assert not app.session_state["show_course_setup"]

    app.button(key="create_course_button").click().run()

    assert not app.exception
    assert app.button(key="create_course_button").label == "Close setup"
    assert app.session_state["show_course_setup"]
    assert any(
        item.key == "setup_course_name" for item in app.text_input
    )

    app.button(key="create_course_button").click().run()

    assert not app.exception
    assert app.button(key="create_course_button").label == "Create course"
    assert not app.session_state["show_course_setup"]
    assert not any(
        item.key == "setup_course_name" for item in app.text_input
    )


def test_dashboard_extracts_txt_and_shows_limits_preview_and_source(monkeypatch):
    upload = install_upload_mock(monkeypatch)
    app = AppTest.from_file(str(DASHBOARD_PATH)).run()
    app.button(key="create_course_button").click().run()
    app.text_input(key="setup_course_name").set_value("TXT course")
    upload["file"] = make_upload(
        "notes.txt",
        "A local TXT fixture.".encode("utf-8"),
    )
    app.run()
    click_extract(app)

    assert not app.exception
    assert app.session_state["setup_document"]["text"] == (
        "A local TXT fixture."
    )
    assert app.session_state["setup_document"]["sources"] == [
        {"filename": "notes.txt", "page": None}
    ]
    assert any("50,000 characters" in item.value for item in app.caption)
    assert any(
        "UTF-8 TXT" in item.value
        and "text-based PDF or PPTX" in item.value
        for item in app.caption
    )
    assert any(
        "notes.txt · TXT file (no page number)" in item.value
        for item in app.caption
    )
    assert "A local TXT fixture." in [item.value for item in app.text]


def test_dashboard_extracts_pdf_and_displays_actual_page_sources(monkeypatch):
    upload = install_upload_mock(monkeypatch)
    app = AppTest.from_file(str(DASHBOARD_PATH)).run()
    app.button(key="create_course_button").click().run()
    pdf_bytes = make_pdf("Page one content.")
    upload["file"] = make_upload("course.pdf", pdf_bytes)
    app.run()
    click_extract(app)

    assert not app.exception
    document = app.session_state["setup_document"]
    assert document["sources"] == [{"filename": "course.pdf", "page": 1}]
    assert document["pages"][0]["page"] == 1
    assert any("course.pdf · Page 1" in item.value for item in app.caption)
    assert any("PDF pages (1)" in item.label for item in app.expander)


def test_dashboard_image_ocr_uses_image_labels_and_updated_copy(monkeypatch):
    upload = install_upload_mock(monkeypatch)
    monkeypatch.setattr(
        "services.document_service._run_ocr", lambda image: "OCR text"
    )
    app = AppTest.from_file(str(DASHBOARD_PATH)).run()
    app.button(key="create_course_button").click().run()
    stream = BytesIO()
    Image.new("RGB", (50, 50), "white").save(stream, format="PNG")
    upload["file"] = make_upload("scan.png", stream.getvalue())
    app.run()
    click_extract(app)

    assert not app.exception
    assert any("Image (1)" in item.label for item in app.expander)
    assert not any("PDF pages" in item.label for item in app.expander)
    assert any("scan.png · Image 1" in item.value for item in app.caption)
    assert any(
        "scanned pdf ocr" in item.value.lower()
        for item in app.caption
    )


def test_dashboard_failure_message_clears_old_document_and_disables_generation(
    monkeypatch,
):
    upload = install_upload_mock(monkeypatch)
    app = AppTest.from_file(str(DASHBOARD_PATH)).run()
    app.button(key="create_course_button").click().run()
    upload["file"] = make_upload("old.txt", b"old content")
    app.run()
    click_extract(app)
    assert app.session_state["setup_document"]["filename"] == "old.txt"

    upload["file"] = make_upload("invalid.txt", b"\xff")
    app.run()

    assert "setup_document" not in app.session_state
    click_extract(app)

    assert not app.exception
    assert "setup_document" not in app.session_state
    assert any(
        "文件必须使用有效的 UTF-8 编码。" in item.value
        for item in app.error
    )
    assert not any(button.label == "Generate topics" for button in app.button)


def test_dashboard_file_replacement_and_removal_never_reuse_old_extraction(
    monkeypatch,
):
    upload = install_upload_mock(monkeypatch)
    app = AppTest.from_file(str(DASHBOARD_PATH)).run()
    app.button(key="create_course_button").click().run()
    upload["file"] = make_upload("first.txt", b"First file.")
    app.run()
    click_extract(app)
    assert app.session_state["setup_document"]["filename"] == "first.txt"

    upload["file"] = make_upload("second.txt", b"Second file.")
    app.run()
    assert "setup_document" not in app.session_state
    assert not any(button.label == "Generate topics" for button in app.button)

    click_extract(app)
    assert app.session_state["setup_document"]["filename"] == "second.txt"

    upload["file"] = None
    app.run()
    assert "setup_document" not in app.session_state
    assert not any(button.label == "Generate topics" for button in app.button)


STUDY_PLAN_PATH = Path(__file__).resolve().parents[1] / "pages" / "1_Study_Hub.py"


def make_study_course(course_id: str, course_name: str) -> StudyPackage:
    return StudyPackage(
        course_id=course_id,
        course_name=course_name,
        topics=[
            {
                "id": f"{course_id}-topic",
                "name": "Course topic",
                "summary": "A short course summary.",
            }
        ],
        relationships=[],
        questions=[],
    )


def test_question_generation_is_disabled_when_course_has_exam_history():
    course = make_study_course("history-course", "History course")
    app = AppTest.from_file(str(STUDY_PLAN_PATH))
    app.session_state["current_course"] = course
    app.session_state["confirmed_scope_history-course"] = [
        "history-course-topic"
    ]
    app.session_state["course_document_history-course"] = {
        "text": "Course material",
    }
    app.session_state["practice_history_history-course"] = {
        "old-session": {"session": object()},
    }
    app.run()

    assert not app.exception
    assert app.button(key="generate_questions_history-course").disabled
    assert any(
        "Question generation is disabled" in item.value
        for item in app.warning
    )




def test_dashboard_extracts_pptx_and_displays_actual_slide_sources(monkeypatch):
    from io import BytesIO

    from pptx import Presentation

    upload = install_upload_mock(monkeypatch)
    app = AppTest.from_file(str(DASHBOARD_PATH)).run()
    app.button(key="create_course_button").click().run()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(0, 0, 2_000_000, 500_000)
    shape.text_frame.text = "PPTX extraction fixture."
    output = BytesIO()
    presentation.save(output)
    upload["file"] = make_upload("slides.pptx", output.getvalue())
    app.run()

    click_extract(app)

    assert not app.exception
    document = app.session_state["setup_document"]
    assert document["sources"] == [{"filename": "slides.pptx", "page": 1}]
    assert document["pages"] == [
        {"page": 1, "text": "PPTX extraction fixture."}
    ]
    assert any(
        "slides.pptx · Slide 1" == item.value for item in app.caption
    )
    assert any("PPTX slides (1)" == item.label for item in app.expander)
    assert any("50,000 characters" in item.value for item in app.caption)
    assert any(
        "UTF-8 TXT" in item.value
        and "text-based PDF or PPTX" in item.value
        for item in app.caption
    )


def test_topic_cards_sync_scope_and_require_reconfirmation():
    course = make_study_course("checkbox-course", "Checkbox course")
    app = AppTest.from_file(str(STUDY_PLAN_PATH))
    app.session_state["current_course"] = course
    app.run()
    checkbox_key = "topic_selected_checkbox-course_checkbox-course-topic"
    app.checkbox(key=checkbox_key).check().run()
    assert not app.exception
    assert app.multiselect[0].value == ["checkbox-course-topic"]
    next(button for button in app.button if button.label == "Confirm exam scope").click().run()
    assert app.session_state["confirmed_scope_checkbox-course"] == ["checkbox-course-topic"]
    app.checkbox(key=checkbox_key).uncheck().run()
    assert not app.exception
    assert app.multiselect[0].value == []
    assert "confirmed_scope_checkbox-course" not in app.session_state
    app.multiselect[0].set_value(["checkbox-course-topic"]).run()
    assert not app.exception
    assert app.checkbox(key=checkbox_key).value
