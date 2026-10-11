from io import BytesIO

import pytest
import pymupdf
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement

from services import document_service
from services.document_service import (
    MAX_FILE_BYTES,
    MAX_TEXT_CHARS,
    extract_pdf,
    extract_pptx,
    extract_txt,
)


def make_pdf(page_texts, *, user_password=None):
    document = pymupdf.open()
    for page_text in page_texts:
        page = document.new_page()
        if page_text:
            page.insert_text((72, 72), page_text)

    options = {}
    if user_password is not None:
        options = {
            "encryption": pymupdf.PDF_ENCRYPT_AES_256,
            "owner_pw": "test-owner-password",
            "user_pw": user_password,
        }

    file_bytes = document.tobytes(**options)
    document.close()
    return file_bytes


def add_grouped_text(slide, text):
    shape = slide.shapes.add_textbox(0, 0, 1_000_000, 300_000)
    shape.text_frame.text = text
    shape_element = shape._element
    shape_parent = shape_element.getparent()

    group = OxmlElement("p:grpSp")
    nonvisual = OxmlElement("p:nvGrpSpPr")
    nonvisual_shape = OxmlElement("p:cNvPr")
    nonvisual_shape.set("id", "99")
    nonvisual_shape.set("name", "Test Group")
    nonvisual.append(nonvisual_shape)
    nonvisual.append(OxmlElement("p:cNvGrpSpPr"))
    nonvisual.append(OxmlElement("p:nvPr"))
    group.append(nonvisual)

    group_properties = OxmlElement("p:grpSpPr")
    transform = OxmlElement("a:xfrm")
    for tag, attributes in (
        ("a:off", {"x": "0", "y": "0"}),
        ("a:ext", {"cx": "1000000", "cy": "300000"}),
        ("a:chOff", {"x": "0", "y": "0"}),
        ("a:chExt", {"cx": "1000000", "cy": "300000"}),
    ):
        element = OxmlElement(tag)
        for key, value in attributes.items():
            element.set(key, value)
        transform.append(element)
    group_properties.append(transform)
    group.append(group_properties)

    shape_parent.remove(shape_element)
    group.append(shape_element)
    slide.shapes._spTree.insert_element_before(group, "p:extLst")


def make_pptx(slide_specs):
    presentation = Presentation()
    for spec in slide_specs:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        for item in spec:
            if item[0] == "text":
                shape = slide.shapes.add_textbox(0, 0, 1_000_000, 300_000)
                shape.text_frame.text = item[1]
            elif item[0] == "table":
                table = slide.shapes.add_table(
                    len(item[1]),
                    len(item[1][0]),
                    0,
                    0,
                    2_000_000,
                    1_000_000,
                ).table
                for row_index, row in enumerate(item[1]):
                    for column_index, value in enumerate(row):
                        table.cell(row_index, column_index).text = value
            elif item[0] == "group":
                add_grouped_text(slide, item[1])

    stream = BytesIO()
    presentation.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize(
    ("text", "filename"),
    [
        ("Plain UTF-8 text", "notes.txt"),
        ("中文内容", "课程资料.txt"),
        ("BOM-prefixed text", "bom.txt"),
        ("Uppercase extension", "NOTES.TXT"),
    ],
)
def test_extract_txt_success_and_page_is_none(text, filename):
    file_bytes = text.encode("utf-8-sig") if filename == "bom.txt" else text.encode("utf-8")

    result = extract_txt(file_bytes, filename)

    assert result == {
        "ok": True,
        "filename": filename,
        "text": text,
        "sources": [{"filename": filename, "page": None}],
    }


def test_rejects_non_txt_extension():
    result = extract_txt(b"text", "notes.pdf")

    assert result["ok"] is False
    assert result["error_code"] == "UNSUPPORTED_FILE_TYPE"


@pytest.mark.parametrize("file_bytes", [b"", b" \t\r\n"])
def test_rejects_empty_content(file_bytes):
    result = extract_txt(file_bytes, "empty.txt")

    assert result["ok"] is False
    assert result["error_code"] == "EMPTY_DOCUMENT"


def test_rejects_invalid_utf8():
    result = extract_txt(b"valid prefix\xffinvalid", "invalid.txt")

    assert result["ok"] is False
    assert result["error_code"] == "INVALID_ENCODING"


def test_file_size_limit_accepts_exact_boundary_before_text_limit():
    file_bytes = b"a" * MAX_FILE_BYTES

    result = extract_txt(file_bytes, "at-size-limit.txt")

    assert result["ok"] is False
    assert result["error_code"] == "TEXT_TOO_LONG"


def test_rejects_file_over_size_limit():
    file_bytes = b"a" * (MAX_FILE_BYTES + 1)

    result = extract_txt(file_bytes, "over-size-limit.txt")

    assert result["ok"] is False
    assert result["error_code"] == "FILE_TOO_LARGE"


def test_text_length_limit_accepts_exact_boundary():
    text = "a" * MAX_TEXT_CHARS

    result = extract_txt(text.encode("utf-8"), "at-text-limit.txt")

    assert result["ok"] is True
    assert result["text"] == text
    assert result["sources"] == [{"filename": "at-text-limit.txt", "page": None}]


def test_rejects_text_over_length_limit():
    text = "a" * (MAX_TEXT_CHARS + 1)

    result = extract_txt(text.encode("utf-8"), "over-text-limit.txt")

    assert result["ok"] is False
    assert result["error_code"] == "TEXT_TOO_LONG"


def test_extract_pdf_returns_page_text_and_real_page_sources():
    file_bytes = make_pdf(["First page", "Second page"])

    result = extract_pdf(file_bytes, "course.pdf")

    assert result["ok"] is True
    assert result["filename"] == "course.pdf"
    assert [page["page"] for page in result["pages"]] == [1, 2]
    assert [page["text"].strip() for page in result["pages"]] == [
        "First page",
        "Second page",
    ]
    assert result["text"] == "\n".join(page["text"] for page in result["pages"])
    assert result["sources"] == [
        {"filename": "course.pdf", "page": 1},
        {"filename": "course.pdf", "page": 2},
    ]


def test_extract_pdf_retains_blank_pages_but_omits_them_from_sources():
    file_bytes = make_pdf(["First page", "", "Third page"])

    result = extract_pdf(file_bytes, "mixed.pdf")

    assert result["ok"] is True
    assert [page["page"] for page in result["pages"]] == [1, 2, 3]
    assert result["pages"][1]["text"] == ""
    assert result["sources"] == [
        {"filename": "mixed.pdf", "page": 1},
        {"filename": "mixed.pdf", "page": 3},
    ]


def test_extract_pdf_accepts_uppercase_extension():
    result = extract_pdf(make_pdf(["Text"]), "COURSE.PDF")

    assert result["ok"] is True
    assert result["filename"] == "COURSE.PDF"


def test_extract_pdf_rejects_pdf_without_extractable_text():
    result = extract_pdf(make_pdf(["", " \t "]), "blank.pdf")

    assert result["ok"] is False
    assert result["error_code"] == "NO_EXTRACTABLE_TEXT"
    assert "OCR" in result["message"]


def test_extract_pdf_rejects_corrupt_content_without_exposing_exception():
    result = extract_pdf(b"not a PDF document", "corrupt.pdf")

    assert result["ok"] is False
    assert result["error_code"] == "INVALID_PDF"
    assert "not a PDF document" not in result["message"]


def test_extract_pdf_rejects_empty_bytes():
    result = extract_pdf(b"", "empty.pdf")

    assert result["ok"] is False
    assert result["error_code"] == "EMPTY_DOCUMENT"


def test_extract_pdf_rejects_encrypted_document():
    file_bytes = make_pdf(["Secret text"], user_password="test-user-password")

    result = extract_pdf(file_bytes, "encrypted.pdf")

    assert result["ok"] is False
    assert result["error_code"] == "ENCRYPTED_PDF"


def test_extract_pdf_rejects_unsupported_file_type():
    result = extract_pdf(make_pdf(["Text"]), "course.txt")

    assert result["ok"] is False
    assert result["error_code"] == "UNSUPPORTED_FILE_TYPE"


def test_extract_pdf_file_size_limit(monkeypatch):
    file_bytes = make_pdf(["Small PDF"])
    monkeypatch.setattr(document_service, "MAX_FILE_BYTES", len(file_bytes))

    at_limit_result = extract_pdf(file_bytes, "at-limit.pdf")
    over_limit_result = extract_pdf(file_bytes + b"x", "over-limit.pdf")

    assert at_limit_result["ok"] is True
    assert over_limit_result["ok"] is False
    assert over_limit_result["error_code"] == "FILE_TOO_LARGE"


def test_extract_pdf_text_length_limit(monkeypatch):
    file_bytes = make_pdf(["Text length boundary"])
    initial_result = extract_pdf(file_bytes, "text-limit.pdf")
    assert initial_result["ok"] is True
    monkeypatch.setattr(
        document_service,
        "MAX_TEXT_CHARS",
        len(initial_result["text"]),
    )

    at_limit_result = extract_pdf(file_bytes, "text-limit.pdf")
    monkeypatch.setattr(
        document_service,
        "MAX_TEXT_CHARS",
        len(initial_result["text"]) - 1,
    )
    over_limit_result = extract_pdf(file_bytes, "text-limit.pdf")

    assert at_limit_result["ok"] is True
    assert over_limit_result["ok"] is False
    assert over_limit_result["error_code"] == "TEXT_TOO_LONG"


def test_extract_pptx_extracts_text_table_group_and_blank_slides():
    file_bytes = make_pptx(
        [
            [
                ("text", "Slide title"),
                ("table", [["Name", "Value"], ["count", "3"]]),
                ("group", "Grouped text"),
            ],
            [],
            [("text", "Final slide")],
        ]
    )

    result = extract_pptx(file_bytes, "course.pptx")

    assert result["ok"] is True
    assert result["filename"] == "course.pptx"
    assert [page["page"] for page in result["pages"]] == [1, 2, 3]
    assert "Slide title" in result["pages"][0]["text"]
    assert "Name\tValue" in result["pages"][0]["text"]
    assert "count\t3" in result["pages"][0]["text"]
    assert "Grouped text" in result["pages"][0]["text"]
    assert result["pages"][1]["text"] == ""
    assert result["pages"][2]["text"] == "Final slide"
    assert result["text"] == "\n".join(page["text"] for page in result["pages"])
    assert result["sources"] == [
        {"filename": "course.pptx", "page": 1},
        {"filename": "course.pptx", "page": 3},
    ]


def test_extract_pptx_accepts_uppercase_extension_and_dispatches():
    file_bytes = make_pptx([[('text', 'Uppercase extension')]])

    result = document_service.extract_document(file_bytes, "COURSE.PPTX")

    assert result["ok"] is True
    assert result["filename"] == "COURSE.PPTX"
    assert result["sources"] == [{"filename": "COURSE.PPTX", "page": 1}]


@pytest.mark.parametrize("file_bytes", [b"", b"not a presentation"])
def test_extract_pptx_rejects_empty_or_corrupt_content(file_bytes):
    result = extract_pptx(file_bytes, "invalid.pptx")

    assert result["ok"] is False
    assert result["error_code"] == (
        "EMPTY_DOCUMENT" if not file_bytes else "INVALID_PPTX"
    )
    assert "not a presentation" not in result["message"]


def test_extract_pptx_rejects_presentation_without_text():
    result = extract_pptx(make_pptx([[], []]), "blank.pptx")

    assert result["ok"] is False
    assert result["error_code"] == "NO_EXTRACTABLE_TEXT"
    assert "OCR" in result["message"]


def test_extract_pptx_rejects_unsupported_extension():
    result = extract_pptx(make_pptx([[('text', 'Text')]]), "slides.ppt")

    assert result["ok"] is False
    assert result["error_code"] == "UNSUPPORTED_FILE_TYPE"


def test_extract_pptx_file_size_limit(monkeypatch):
    file_bytes = make_pptx([[('text', 'File size limit')]])
    monkeypatch.setattr(document_service, "MAX_FILE_BYTES", len(file_bytes))

    at_limit_result = extract_pptx(file_bytes, "at-limit.pptx")
    over_limit_result = extract_pptx(file_bytes + b"x", "over-limit.pptx")

    assert at_limit_result["ok"] is True
    assert over_limit_result["ok"] is False
    assert over_limit_result["error_code"] == "FILE_TOO_LARGE"


def test_extract_pptx_text_length_limit(monkeypatch):
    file_bytes = make_pptx([[('text', 'PPTX length boundary')]])
    initial_result = extract_pptx(file_bytes, "text-limit.pptx")
    assert initial_result["ok"] is True
    text_length = len(initial_result["text"])
    monkeypatch.setattr(document_service, "MAX_TEXT_CHARS", text_length)

    at_limit_result = extract_pptx(file_bytes, "text-limit.pptx")
    monkeypatch.setattr(document_service, "MAX_TEXT_CHARS", text_length - 1)
    over_limit_result = extract_pptx(file_bytes, "text-limit.pptx")

    assert at_limit_result["ok"] is True
    assert over_limit_result["ok"] is False
    assert over_limit_result["error_code"] == "TEXT_TOO_LONG"


@pytest.mark.parametrize(
    ("filename", "selected_extractor"),
    [
        ("notes.txt", "txt"),
        ("NOTES.TXT", "txt"),
        ("course.pdf", "pdf"),
        ("COURSE.PDF", "pdf"),
    ],
)
@pytest.mark.parametrize("ok", [True, False])
def test_extract_document_dispatches_and_returns_result_unchanged(
    monkeypatch,
    filename,
    selected_extractor,
    ok,
):
    file_bytes = b"in-memory file bytes"
    result = {"ok": ok, "custom_field": object()}
    calls = []

    def fake_txt(received_bytes, received_filename):
        calls.append(("txt", received_bytes, received_filename))
        return result

    def fake_pdf(received_bytes, received_filename):
        calls.append(("pdf", received_bytes, received_filename))
        return result

    monkeypatch.setattr(document_service, "extract_txt", fake_txt)
    monkeypatch.setattr(document_service, "extract_pdf", fake_pdf)

    extracted = document_service.extract_document(file_bytes, filename)

    assert extracted is result
    assert calls == [(selected_extractor, file_bytes, filename)]


def test_extract_document_rejects_unsupported_type_without_calling_extractors(
    monkeypatch,
):
    def fail_if_called(*args):
        pytest.fail("An extractor must not be called for unsupported file types")

    monkeypatch.setattr(document_service, "extract_txt", fail_if_called)
    monkeypatch.setattr(document_service, "extract_pdf", fail_if_called)

    result = document_service.extract_document(b"data", "notes.docx")

    assert result == {
        "ok": False,
        "error_code": "UNSUPPORTED_FILE_TYPE",
        "message": "仅支持 TXT、文字型 PDF 或文字型 PPTX。",
    }
