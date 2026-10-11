from io import BytesIO

import pymupdf
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_TEXT_CHARS = 50_000


def _shape_texts(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shape_texts(shape.shapes)
        elif shape.has_table:
            for row in shape.table.rows:
                yield "\t".join(cell.text for cell in row.cells)
        elif shape.has_text_frame:
            yield shape.text_frame.text


def extract_txt(file_bytes: bytes, filename: str) -> dict:
    if not filename.lower().endswith(".txt"):
        return {
            "ok": False,
            "error_code": "UNSUPPORTED_FILE_TYPE",
            "message": "仅支持 .txt 文件。",
        }

    if len(file_bytes) > MAX_FILE_BYTES:
        return {
            "ok": False,
            "error_code": "FILE_TOO_LARGE",
            "message": "文件大小不能超过 5 MiB。",
        }

    try:
        text = file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        return {
            "ok": False,
            "error_code": "INVALID_ENCODING",
            "message": "文件必须使用有效的 UTF-8 编码。",
        }

    if not text.strip():
        return {
            "ok": False,
            "error_code": "EMPTY_DOCUMENT",
            "message": "文件内容为空。",
        }

    if len(text) > MAX_TEXT_CHARS:
        return {
            "ok": False,
            "error_code": "TEXT_TOO_LONG",
            "message": "文本长度不能超过 50,000 个字符。",
        }

    return {
        "ok": True,
        "filename": filename,
        "text": text,
        "sources": [{"filename": filename, "page": None}],
    }


def extract_pdf(file_bytes: bytes, filename: str) -> dict:
    if not filename.lower().endswith(".pdf"):
        return {
            "ok": False,
            "error_code": "UNSUPPORTED_FILE_TYPE",
            "message": "仅支持 .pdf 文件。",
        }

    if not file_bytes:
        return {
            "ok": False,
            "error_code": "EMPTY_DOCUMENT",
            "message": "PDF 文件内容为空。",
        }

    if len(file_bytes) > MAX_FILE_BYTES:
        return {
            "ok": False,
            "error_code": "FILE_TOO_LARGE",
            "message": "文件大小不能超过 5 MiB。",
        }

    try:
        with pymupdf.open(stream=file_bytes, filetype="pdf") as document:
            if document.needs_pass:
                return {
                    "ok": False,
                    "error_code": "ENCRYPTED_PDF",
                    "message": "PDF 需要密码才能读取，暂不支持加密 PDF。",
                }

            pages = []
            sources = []
            total_text_chars = 0

            for page_number, page in enumerate(document, start=1):
                page_text = page.get_text()
                pages.append({"page": page_number, "text": page_text})

                if page_number > 1:
                    total_text_chars += 1
                total_text_chars += len(page_text)
                if total_text_chars > MAX_TEXT_CHARS:
                    return {
                        "ok": False,
                        "error_code": "TEXT_TOO_LONG",
                        "message": "PDF 提取文本长度超过 50,000 个字符。",
                    }

                if page_text.strip():
                    sources.append({"filename": filename, "page": page_number})

            if not sources:
                return {
                    "ok": False,
                    "error_code": "NO_EXTRACTABLE_TEXT",
                    "message": "PDF 中没有可提取的文字，可能是扫描件或空白 PDF；本功能不支持 OCR。",
                }

            text = "\n".join(page["text"] for page in pages)
    except Exception:
        return {
            "ok": False,
            "error_code": "INVALID_PDF",
            "message": "无法读取 PDF，文件可能已损坏或内容不是有效的 PDF。",
        }

    return {
        "ok": True,
        "filename": filename,
        "text": text,
        "pages": pages,
        "sources": sources,
    }


def extract_pptx(file_bytes: bytes, filename: str) -> dict:
    if not filename.lower().endswith(".pptx"):
        return {
            "ok": False,
            "error_code": "UNSUPPORTED_FILE_TYPE",
            "message": "仅支持 .pptx 文件。",
        }

    if not file_bytes:
        return {
            "ok": False,
            "error_code": "EMPTY_DOCUMENT",
            "message": "PPTX 文件内容为空。",
        }

    if len(file_bytes) > MAX_FILE_BYTES:
        return {
            "ok": False,
            "error_code": "FILE_TOO_LARGE",
            "message": "文件大小不能超过 5 MiB。",
        }

    try:
        presentation = Presentation(BytesIO(file_bytes))
        pages = []
        sources = []
        total_text_chars = 0

        for slide_number, slide in enumerate(presentation.slides, start=1):
            slide_text = "\n".join(
                text for text in _shape_texts(slide.shapes) if text
            )
            pages.append({"page": slide_number, "text": slide_text})

            if slide_number > 1:
                total_text_chars += 1
            total_text_chars += len(slide_text)
            if total_text_chars > MAX_TEXT_CHARS:
                return {
                    "ok": False,
                    "error_code": "TEXT_TOO_LONG",
                    "message": "PPTX 提取文本长度超过 50,000 个字符。",
                }

            if slide_text.strip():
                sources.append({"filename": filename, "page": slide_number})

        if not sources:
            return {
                "ok": False,
                "error_code": "NO_EXTRACTABLE_TEXT",
                "message": "PPTX 中没有可提取的文字，可能是空白幻灯片或图片内容；本功能不支持 OCR。",
            }

        text = "\n".join(slide["text"] for slide in pages)
    except Exception:
        return {
            "ok": False,
            "error_code": "INVALID_PPTX",
            "message": "无法读取 PPTX，文件可能已损坏或内容不是有效的 PPTX。",
        }

    return {
        "ok": True,
        "filename": filename,
        "text": text,
        "pages": pages,
        "sources": sources,
    }


def extract_document(file_bytes: bytes, filename: str) -> dict:
    if filename.lower().endswith(".txt"):
        return extract_txt(file_bytes, filename)

    if filename.lower().endswith(".pdf"):
        return extract_pdf(file_bytes, filename)

    if filename.lower().endswith(".pptx"):
        return extract_pptx(file_bytes, filename)

    return {
        "ok": False,
        "error_code": "UNSUPPORTED_FILE_TYPE",
        "message": "仅支持 TXT、文字型 PDF 或文字型 PPTX。",
    }
